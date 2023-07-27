import tkinter as tk
from tkinter import filedialog
from openpyxl import load_workbook

def load_excel_file():
    file_path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx *.xlsm")])
    if file_path:
        entry.delete(0, tk.END)
        entry.insert(0, file_path)
        workbook = load_workbook(file_path)
        sheet_names = workbook.sheetnames

        for idx, sheet_name in enumerate(sheet_names):
            tk.Radiobutton(root, text=sheet_name, variable=sheet_var, value=idx).pack(anchor=tk.W)

        btn_select.config(state=tk.NORMAL)

def select_sheet():
    selected_sheet_index = sheet_var.get()
    selected_sheet_name = workbook.sheetnames[selected_sheet_index]
    print(f"Selected sheet: {selected_sheet_name}")

# Create the main window
root = tk.Tk()
root.title("Excel Sheet Selector")

# Create a variable to hold the selected sheet index
sheet_var = tk.IntVar()

# Create the input entry and load button
entry = tk.Entry(root, width=50)
entry.pack(pady=10)
btn_load = tk.Button(root, text="Load Excel File", command=load_excel_file)
btn_load.pack(pady=5)

# Create a button to select the sheet
btn_select = tk.Button(root, text="Select Sheet", command=select_sheet, state=tk.DISABLED)
btn_select.pack(pady=5)

root.mainloop()
